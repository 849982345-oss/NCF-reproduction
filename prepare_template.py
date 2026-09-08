#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Prepare working template: strip 'INNER CHAPTER' master watermark."""
import shutil
from pptx import Presentation

SRC = r"C:\Users\Administrator\AppData\Local\hermes\skills\shared\pptx-from-layouts\templates\inner-chapter.pptx"
DST = r"C:\research-training\repro-ncf\ncf_inner_template.pptx"
shutil.copy(SRC, DST)
p = Presentation(DST)
removed = 0
for master in p.slide_masters:
    for sh in list(master.shapes):
        if sh.has_text_frame and "INNER" in sh.text_frame.text.upper():
            sh._element.getparent().remove(sh._element)
            removed += 1
p.save(DST)
print("removed", removed, "watermarks ->", DST)
