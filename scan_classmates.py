#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Outline extraction from classmates' PPTs for style/structure reference."""
import os, glob
from pptx import Presentation

SRC = r"C:\Users\Administrator\Desktop\研一汇报PPT"
OUT = r"C:\research-training\repro-ncf\classmates_outline.txt"

def first_text(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().replace("\n", " / ")[:70]
    return "(no text)"

lines = []
for f in sorted(glob.glob(os.path.join(SRC, "*.pptx"))):
    name = os.path.basename(f)
    try:
        prs = Presentation(f)
        lines.append(f"\n{'='*80}\n### {name}  |  slides: {len(prs.slides.__iter__.__self__._sldIdLst)}\n")
        for i, s in enumerate(prs.slides, 1):
            texts = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    texts.append(sh.text_frame.text.strip())
                elif sh.has_table:
                    rows = len(sh.table.rows); cols = len(sh.table.columns)
                    texts.append(f"[TABLE {rows}x{cols}]")
                elif sh.shape_type == 13:
                    texts.append("[IMAGE]")
            body = " || ".join(t.replace("\n", " / ")[:110] for t in texts[:6])
            lines.append(f"  S{i:02d}: {body[:400]}")
    except Exception as e:
        lines.append(f"\n### {name}  ERROR: {e}")

open(OUT, "w", encoding="utf-8").write("\n".join(lines))
print("saved", OUT, "chars:", sum(len(l) for l in lines))
