#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Probe placeholders of key inner-chapter layouts."""
from pptx import Presentation
from pptx.util import Emu

TPL = r"C:\Users\Administrator\AppData\Local\hermes\skills\shared\pptx-from-layouts\templates\inner-chapter.pptx"
p = Presentation(TPL)
want = ["title-cover", "column-3-centered-a", "title-centered", "content-image-right-a", "content-centered-a", "column-2-centered"]
for lay in p.slide_layouts:
    if lay.name not in want:
        continue
    print(f"\n== layout: {lay.name} ==")
    for ph in lay.placeholders:
        try:
            pos = f"L{Emu(ph.left).inches:.2f} T{Emu(ph.top).inches:.2f} W{Emu(ph.width).inches:.2f} H{Emu(ph.height).inches:.2f}"
        except Exception:
            pos = "no-pos"
        print(f"  idx={ph.placeholder_format.idx} type={ph.placeholder_format.type} name='{ph.name}' {pos}")
